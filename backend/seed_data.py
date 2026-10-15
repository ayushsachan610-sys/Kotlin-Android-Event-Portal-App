import os
import sys
from pathlib import Path

# Ensure UTF-8 output on Windows consoles
if sys.stdout and hasattr(sys.stdout, "reconfigure"):
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass
import openpyxl
from backend.database import get_db_connection, init_db

DATA_FILE_PATH = Path(__file__).resolve().parent.parent / "data" / "events_seed.xlsx"


def seed_events():
    """Seeds the events table from data/events_seed.xlsx if the table is currently empty."""
    # Ensure database tables exist first
    init_db()

    # Verify seed file existence
    if not DATA_FILE_PATH.is_file():
        raise FileNotFoundError(
            f"The seed file '{DATA_FILE_PATH}' is missing. "
            f"Please copy the provided events_seed.xlsx file into the data/ folder before running this script."
        )

    conn = get_db_connection()
    cursor = conn.cursor()

    # Check if events already exist
    cursor.execute("SELECT COUNT(*) AS count FROM events")
    count = cursor.fetchone()["count"]
    if count > 0:
        print(f"[Seed] events table already contains {count} rows. Skipping seed to prevent duplication.")
        conn.close()
        return

    # Load workbook
    wb = openpyxl.load_workbook(DATA_FILE_PATH, data_only=True)
    if "Events" not in wb.sheetnames:
        raise ValueError(f"Sheet 'Events' not found in {DATA_FILE_PATH}. Available sheets: {wb.sheetnames}")

    sheet = wb["Events"]

    # Verify header
    expected_header = ["id", "title", "category", "description", "event_date", "event_time", "venue", "capacity", "fee", "icon"]
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        raise ValueError("Seed spreadsheet is empty.")

    header = [str(col).strip() if col is not None else "" for col in rows[0]]
    if header != expected_header:
        print(f"[Warning] Header row {header} does not match exact expected order {expected_header}. Parsing by column index.")

    inserted_count = 0
    for row in rows[1:]:
        if not row or all(v is None for v in row):
            continue
        
        # Columns in order: id, title, category, description, event_date, event_time, venue, capacity, fee, icon
        # We can either let SQLite assign autoincrement ID or explicitly preserve the ID from the sheet
        event_id, title, category, description, event_date, event_time, venue, capacity, fee, icon = row[:10]

        cursor.execute("""
            INSERT INTO events (id, title, category, description, event_date, event_time, venue, capacity, fee, icon)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """, (
            int(event_id) if event_id is not None else None,
            str(title).strip() if title is not None else "",
            str(category).strip() if category is not None else "",
            str(description).strip() if description is not None else "",
            str(event_date).strip() if event_date is not None else "",
            str(event_time).strip() if event_time is not None else "",
            str(venue).strip() if venue is not None else "",
            int(capacity) if capacity is not None else 0,
            int(fee) if fee is not None else 0,
            str(icon).strip() if icon is not None else ""
        ))
        inserted_count += 1

    conn.commit()
    conn.close()
    print(f"[Seed] Successfully seeded {inserted_count} events from {DATA_FILE_PATH}.")


if __name__ == "__main__":
    seed_events()

    # Query and display table contents
    conn = get_db_connection()
    events = conn.execute("SELECT id, title, category, event_date, event_time, venue, capacity, fee, icon FROM events").fetchall()
    print("\n--- Current Events in Database ---")
    for ev in events:
        print(f"[{ev['id']}] {ev['icon']} {ev['title']} | {ev['category']} | {ev['event_date']} {ev['event_time']} | {ev['venue']} | Cap: {ev['capacity']} | Fee: ₹{ev['fee']}")
    conn.close()
